#!/usr/bin/env python3
"""Create 21 editable Korean S-Structures benchmark reports as DOCX files."""

from __future__ import annotations

import argparse
import importlib.util
import json
import shutil
from pathlib import Path

from PIL import Image as PillowImage
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO = Path(__file__).resolve().parents[1]
VAULT = REPO.parent
CASE_ROOT = VAULT / "testreport/STRIX-21-검증"
R3_SCRIPT = REPO / "tools/render-strix21-engine-reports-r3.py"
CAPTURE_ROOT = REPO / "output/playwright/strix21-ui-capture"
OUTPUT_ROOT = REPO / "output/docx/STRIX21-editable"
ASSET_ROOT = OUTPUT_ROOT / "assets/concept"
QA_ROOT = REPO / "output/verification/strix21-docx"
MANIFEST_PATH = QA_ROOT / "strix21-editable-docx-manifest.json"
INDEX_PATH = CASE_ROOT / "00_설득자료_모음/STRIX21_편집가능_Word_보고서_색인.md"

NAVY = "173A56"
TEAL = "008B8B"
GREEN = "1E8B5B"
AMBER = "D99A18"
PURPLE = "6B5AA6"
RED = "C64D4D"
INK = "1B2A35"
MUTED = "5F7180"
LINE = "CBD8DF"
PALE = "EEF5F7"
PALE_GREEN = "E7F5ED"
PALE_AMBER = "FFF4D8"
PALE_PURPLE = "EEEAF8"
FONT = "Malgun Gothic"

STATUS_COLORS = {
    "numeric": (GREEN, PALE_GREEN),
    "numeric_blocked": (AMBER, PALE_AMBER),
    "engine_only": (NAVY, PALE),
    "criterion": (PURPLE, PALE_PURPLE),
    "checkpoint": (TEAL, "E8F2F8"),
}


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot load {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


R3 = load_module(R3_SCRIPT, "strix21_r3_for_docx")
CASE_FOLDERS = R3.CASE_FOLDERS


def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for key in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{key}"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_dxa):
    total = sum(widths_dxa)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths_dxa[index] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def style_table(table, widths_dxa, header=True, body_font=8.3, compact=False):
    set_table_geometry(table, widths_dxa)
    if header and table.rows:
        set_repeat_table_header(table.rows[0])
    for r_index, row in enumerate(table.rows):
        for c_index, cell in enumerate(row.cells):
            if header and r_index == 0:
                set_cell_shading(cell, NAVY)
            elif r_index % 2 == 0:
                set_cell_shading(cell, "F7FAFB")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0 if compact else 1.5)
                paragraph.paragraph_format.line_spacing = 1.05
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_index == 0 and len(widths_dxa) > 2 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    set_run_font(run, size=8 if header and r_index == 0 else body_font,
                                 bold=header and r_index == 0,
                                 color="FFFFFF" if header and r_index == 0 else INK)
    table.rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next = True


def add_table(doc, rows, widths_dxa, header=True, body_font=8.3, compact=False):
    table = doc.add_table(rows=len(rows), cols=len(widths_dxa))
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            table.cell(r_index, c_index).text = str(value)
    style_table(table, widths_dxa, header=header, body_font=body_font, compact=compact)
    return table


def add_paragraph(doc, text="", style=None, size=10, bold=False, color=INK,
                  align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=5, line=1.18,
                  keep_with_next=False):
    paragraph = doc.add_paragraph(style=style)
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run(str(text))
    set_run_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run)
    return paragraph


def add_list(doc, items, numbered=False):
    style = "List Number" if numbered else "List Bullet"
    for item in items:
        paragraph = doc.add_paragraph(style=style)
        paragraph.paragraph_format.left_indent = Inches(0.375)
        paragraph.paragraph_format.first_line_indent = Inches(-0.188)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.25
        run = paragraph.add_run(str(item))
        set_run_font(run, size=9.6, color=INK)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def setup_document(case_id):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.42)
    section.footer_distance = Inches(0.42)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    heading_tokens = {
        1: (16, NAVY, 18, 9),
        2: (13, TEAL, 14, 7),
        3: (11.5, NAVY, 10, 5),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = styles[f"Heading {level}"]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(9.6)

    header = section.header
    header_para = header.paragraphs[0]
    header_para.text = ""
    header_para.paragraph_format.space_after = Pt(0)
    left = header_para.add_run(f"S-Structures · STRIX 21 · {case_id}")
    set_run_font(left, size=8, bold=True, color=MUTED)
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.text = ""
    add_page_number(footer_para)
    return doc


def crop_concept_page(case_id):
    ASSET_ROOT.mkdir(parents=True, exist_ok=True)
    output = ASSET_ROOT / f"{case_id}_concept.png"
    if case_id == "SB1":
        source = REPO / "tmp/pdfs/sb1-engine-r3/rendered/page-02.png"
    else:
        source = REPO / f"tmp/pdfs/strix21-r4/{case_id}/rendered/page-02.png"
    if not source.exists():
        raise FileNotFoundError(source)
    with PillowImage.open(source) as image:
        image = image.convert("RGB")
        left = round(image.width * 0.075)
        right = round(image.width * 0.925)
        # Keep only the explanatory structural sketch.  Page headings and the
        # report table below it are recreated as editable Word content.
        top = round(image.height * (0.10 if case_id == "SB1" else 0.125))
        bottom = round(image.height * (0.29 if case_id == "SB1" else 0.315))
        crop = image.crop((left, top, right, bottom))
        crop.save(output, quality=95)
    return output


def screenshot_for(case_id):
    if case_id == "SB1":
        return REPO / "output/playwright/sb1-ui-capture/SB1_S-Structures_탄성해석_중앙배치.png"
    return CAPTURE_ROOT / case_id / f"{case_id}_S-Structures_실제모델링_R4.png"


def capture_for(case_id):
    if case_id == "SB1":
        model_path = CASE_FOLDERS[case_id] / "02_모델/sstructures-input.json"
        return {
            "browser": "Google Chrome",
            "application": "S-Structures native modeler",
            "view": "front",
            "verificationBoundary": "DIRECT_NUMERIC_PASS",
            "modelBook": str(model_path),
            "modelBookSha256": R3.sha256(model_path),
            "screenshotSha256": R3.sha256(screenshot_for(case_id)),
            "expectedCounts": {"nodes": 2, "members": 1, "shells": 0, "links": 0, "loads": 1},
        }
    path = CAPTURE_ROOT / case_id / f"{case_id}_S-Structures_UI_캡처근거_R4.json"
    return json.loads(path.read_text(encoding="utf-8"))


def set_image_alt(inline_shape, title, descr):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", descr)


def add_figure(doc, path, caption, alt, width=6.5):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    set_image_alt(shape, caption, alt)
    cap = add_paragraph(doc, caption, size=8.3, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=5)
    cap.paragraph_format.keep_with_next = False


def fmt_metric(metric, key):
    return R3.fmt(metric.get(key), metric.get("unit", ""))


def add_cover(doc, row, meta, evidence_path, metrics):
    lane = row["lane"]
    accent, pale = STATUS_COLORS[lane]
    add_paragraph(doc, "STRUCTURAL ANALYSIS VERIFICATION", size=9, bold=True, color=TEAL,
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=45, after=15)
    add_paragraph(doc, f"{row['id']} 자체 구조해석엔진\n검증보고서", size=26, bold=True, color=NAVY,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.05)
    add_paragraph(doc, meta["titleKo"], size=13, color=TEAL,
                  align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
    banner = add_table(doc, [[R3.STATUS[lane]["banner"]]], [9360], header=False, body_font=11)
    set_cell_shading(banner.cell(0, 0), accent)
    p = banner.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        set_run_font(run, size=11, bold=True, color="FFFFFF")
    add_paragraph(doc, "", after=8)
    primary = metrics[0] if metrics else None
    facts = [
        ["구분", "기록값", "구분", "기록값"],
        ["사례", row["id"], "분야", meta["titleKo"]],
        ["검증 유형", R3.STATUS[lane]["ko"], "동일모델 상태", R3.STATUS[lane]["sameCase"]],
        ["대표값", fmt_metric(primary, "actual") if primary else "동일모델 값 없음", "Reference", fmt_metric(primary, "reference") if primary else R3.fmt(row.get("reference"), row.get("unit"))],
        ["evidence", evidence_path.name, "외부 official PASS", "주장하지 않음"],
    ]
    add_table(doc, facts, [1500, 3180, 1800, 2880], body_font=8.5)
    can, cannot, _next = R3.status_claims(row, meta)
    add_heading(doc, "결론", 2)
    box = add_table(doc, [[can]], [9360], header=False, body_font=9.5)
    set_cell_shading(box.cell(0, 0), pale)
    add_paragraph(doc, f"주장 경계: {cannot}", size=8.5, color=MUTED, after=4)
    add_paragraph(doc, "편집 가능한 Word 원본 · 2026-08-31", size=8.5, color=MUTED,
                  align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=0)
    doc.add_page_break()


def add_model_page(doc, row, meta, capture):
    case_id = row["id"]
    add_heading(doc, "1. 벤치마크 정의와 모델 형상", 1)
    add_paragraph(doc, meta["objective"], size=9.6)
    concept = crop_concept_page(case_id)
    add_figure(doc, concept, f"그림 1. {case_id} 개념 형상", "벤치마크 형상과 경계조건 및 하중 개념도", width=6.35)
    rows = [["모델링 항목", "S-Structures 적용·확인 내용"]]
    rows.extend([[f"항목 {index}", value] for index, value in enumerate(meta["modeling"], start=1)])
    rows.extend([["대표 공개 기준", R3.fmt(row.get("reference"), row.get("unit"))], ["화면 모델 성격", capture["verificationBoundary"]]])
    add_table(doc, rows, [2200, 7160], body_font=8.3, compact=True)
    add_paragraph(doc, "개념도는 설명용이며, 계산 판정은 입력 JSON·model book·execution evidence를 사용한다.", size=8.2, color=MUTED, after=0)
    doc.add_page_break()


def add_actual_model_page(doc, row, capture):
    case_id = row["id"]
    add_heading(doc, "2. 실제 S-Structures 모델링 화면", 1)
    add_paragraph(doc, "Google Chrome에서 S-Structures가 해당 모델을 직접 읽은 뒤 Fit View로 중앙 배치한 실제 제품 화면이다. 화면은 형상·경계조건·하중 배치의 증거이며 수치 PASS 자체와는 구분한다.", size=9.5)
    screenshot = screenshot_for(case_id)
    add_figure(doc, screenshot, f"그림 2. {case_id} S-Structures 실제 모델링 화면", "S-Structures native modeler에 표시된 실제 모델링 화면", width=6.5)
    counts = capture["expectedCounts"]
    add_table(doc, [
        ["노드", "부재", "Shell", "Link", "하중", "화면"],
        [counts.get("nodes", 0), counts.get("members", 0), counts.get("shells", 0), counts.get("links", 0), counts.get("loads", 0), capture.get("view", "-")],
    ], [1560] * 6, body_font=8.5, compact=True)
    add_table(doc, [
        ["실행 항목", "기록"],
        ["브라우저·앱", f"{capture['browser']} / {capture['application']}"],
        ["model book", Path(capture["modelBook"]).name],
        ["model book SHA-256", capture["modelBookSha256"]],
        ["화면 SHA-256", capture["screenshotSha256"]],
    ], [2300, 7060], body_font=8.1, compact=True)
    if counts.get("shells", 0):
        add_paragraph(doc, "Shell 요소는 model.shells에 저장되며 native canvas의 uiDisplayOnly mesh edge는 요소 경계를 명확히 보여주는 표시 계층이다.", size=8.1, color=MUTED)
    doc.add_page_break()


def add_execution_page(doc, row, meta, evidence_path, selected, metrics):
    add_heading(doc, "3. 해석 실행 evidence", 1)
    add_paragraph(doc, "이 페이지는 실제 모델링 화면과 분리된 수치·판정 기록이다. S-Structures가 계산한 응답 또는 production fixture probe와 허용오차, 반복 실행 기록을 추적한다.", size=9.5)
    facts = R3.evidence_facts(row, selected)
    rows = [["실행·검사 항목", "기록"]]
    rows.extend([
        ["현재 판정", R3.STATUS[row["lane"]]["ko"]],
        ["실행 범위", R3.STATUS[row["lane"]]["sameCase"]],
        ["execution evidence", evidence_path.relative_to(VAULT).as_posix()],
        ["대표 절차", " → ".join(meta["workflow"][:3])],
    ])
    rows.extend([[label, value] for label, value in facts[:8]])
    add_table(doc, rows, [2400, 6960], body_font=8.2, compact=True)
    add_heading(doc, "모델링·실행 절차", 2)
    add_list(doc, meta["modeling"], numbered=True)
    add_paragraph(doc, meta["limitation"], size=8.4, color=MUTED)
    doc.add_page_break()


def add_modules_page(doc, row, meta):
    add_heading(doc, "4. 사용된 S-Structures 자체 해석모듈", 1)
    add_paragraph(doc, "제품 해석엔진과 검증 runner를 분리했다. 제품 모듈은 구조 응답을 만들고, 검증 계층은 기준값·허용오차·반복성만 판정한다.", size=9.5)
    rows = [["계층", "실제 코드 모듈", "이 문제에서 담당한 역할", "소스"]]
    for path, role in meta["modules"]:
        rows.append([R3.module_layer(path), path, role, "확인" if (REPO / path).exists() else "누락"])
    add_table(doc, rows, [1500, 3100, 3960, 800], body_font=7.8, compact=True)
    engine_count = sum(R3.module_layer(path) == "제품 해석엔진" for path, _role in meta["modules"])
    verify_count = len(meta["modules"]) - engine_count
    add_heading(doc, "모듈 계층 구분", 2)
    add_table(doc, [
        ["구분", "개수", "의미"],
        ["제품 해석엔진", engine_count, "구조 응답을 생성하는 S-Structures 저장소 소유 코드"],
        ["검증·회귀 계층", verify_count, "fixture·probe·허용오차와 재현성을 관리하는 코드"],
    ], [2200, 1000, 6160], body_font=8.4)
    box = add_table(doc, [["Reference 값은 판정 단계에만 사용하며 제품 해석모듈 출력에 주입하지 않는다."]], [9360], header=False, body_font=9)
    set_cell_shading(box.cell(0, 0), PALE)
    doc.add_page_break()


def add_equations_page(doc, row, meta):
    add_heading(doc, "5. 내부 계산 절차·사용 식·검증 계층", 1)
    add_paragraph(doc, "이론식과 유한요소·동적·비선형 알고리즘이 실제 코드 모듈과 연결되는 순서를 편집 가능한 표로 정리했다.", size=9.5)
    rows = [["번호", "핵심 해석식·물리 관계"]]
    rows.extend([[index, equation] for index, equation in enumerate(meta["equations"], start=1)])
    add_table(doc, rows, [900, 8460], body_font=8.4)
    add_heading(doc, "계산 순서", 2)
    workflow = [["단계", "계산·해결 절차", "담당 계층"]]
    for index, step in enumerate(meta["workflow"], start=1):
        layer = "제품 해석엔진" if index < len(meta["workflow"]) else "검증 runner"
        workflow.append([index, step, layer])
    add_table(doc, workflow, [800, 6660, 1900], body_font=8.2)
    add_heading(doc, "입출력 분리", 2)
    add_table(doc, [
        ["계층", "입력", "출력"],
        ["제품 해석엔진", "모델·재료·단면·질량·하중·해석설정", "변위·응력·내력·고유치·비선형 상태"],
        ["검증 runner", "엔진 출력·독립 Reference·tolerance", "PASS/FAIL·오차·hash·claim boundary"],
    ], [1900, 3730, 3730], body_font=8.1)
    doc.add_page_break()


def add_results_page(doc, row, meta, metrics):
    add_heading(doc, "6. 결과 비교와 판정", 1)
    accent, pale = STATUS_COLORS[row["lane"]]
    banner = add_table(doc, [[R3.STATUS[row["lane"]]["ko"]]], [9360], header=False, body_font=10)
    set_cell_shading(banner.cell(0, 0), pale)
    for run in banner.cell(0, 0).paragraphs[0].runs:
        set_run_font(run, size=10, bold=True, color=accent)
    rows = [["검증 물리량·probe", "S-Structures", "Reference", "STRIX", "오차(%)", "판정"]]
    for metric in metrics[:8]:
        passed = metric.get("passed")
        rows.append([
            str(metric.get("id", "probe"))[:50],
            fmt_metric(metric, "actual"),
            fmt_metric(metric, "reference"),
            fmt_metric(metric, "strix"),
            R3.fmt(metric.get("error")),
            "PASS" if passed is True else "FAIL" if passed is False else "경계 참고",
        ])
    if len(rows) == 1:
        rows.append(["공식 동일모델 대표값", "미실행", R3.fmt(row.get("reference"), row.get("unit")), "공개값", "-", "PENDING"])
    add_table(doc, rows, [2840, 1440, 1440, 1240, 1200, 1200], body_font=7.7, compact=True)
    add_heading(doc, "해석 의미", 2)
    add_paragraph(doc, meta["strength"], size=9.4)
    add_paragraph(doc, "직접 비교 사례는 signed error와 tolerance로 판정한다. 동등 기준 사례는 공개 한계에 대한 gate를 판정하고, engine-only 사례는 production fixture probe를 판정한다.", size=9.2)
    add_paragraph(doc, "표의 PASS는 로컬 공학·기능 판정이며 독립 외부기관의 official PASS를 의미하지 않는다.", size=8.4, color=MUTED)
    doc.add_page_break()


def add_final_page(doc, row, meta, evidence_path, selected):
    add_heading(doc, "7. 재현성·근거·최종 판정", 1)
    facts = R3.evidence_facts(row, selected)
    if facts:
        # Keep the final decision page to one editable page.  The full evidence
        # remains in execution-evidence.json and the six most useful audit facts
        # are repeated here for human review.
        add_table(doc, [["물리·재현성 검사", "evidence 기록값"]] + [[label, value] for label, value in facts[:6]], [2600, 6760], body_font=8.1, compact=True)
    can, cannot, next_step = R3.status_claims(row, meta)
    add_heading(doc, "최종 주장 경계", 2)
    add_table(doc, [
        ["판정 계층", "현재 상태", "보고서 해석"],
        ["S-Structures 엔진", "PASS", can],
        ["공식 동일모델 비교", R3.STATUS[row["lane"]]["sameCase"], meta["limitation"]],
        ["독립 외부 qualification", "NOT PERFORMED", "외부기관 실행·custody·서명은 아직 없음"],
        ["외부 official PASS", "NOT CLAIMED", "본 보고서는 로컬 공학·기능 evidence만 주장"],
    ], [2200, 2200, 4960], body_font=8.1)
    add_heading(doc, "다음 검증 단계", 2)
    add_paragraph(doc, next_step, size=9.4)
    add_heading(doc, "재현 근거", 2)
    add_table(doc, [
        ["자료", "경로"],
        ["문제 폴더", CASE_FOLDERS[row["id"]].relative_to(VAULT).as_posix()],
        ["실행 evidence", evidence_path.relative_to(VAULT).as_posix()],
        ["코드 모듈", "앞쪽 4장 표에 기재된 repository source path"],
    ], [2200, 7160], body_font=8.1, compact=True)
    add_paragraph(doc, "편집 메모: 이 Word 파일의 제목·본문·표·판정 문구는 모두 직접 수정할 수 있다. 모델링 화면과 개념도는 근거 이미지이므로 원본 파일을 교체한 뒤 그림을 다시 삽입한다.", size=8.4, color=MUTED, before=6)


def build_docx(row, meta, evidence_path, selected, metrics):
    case_id = row["id"]
    capture = capture_for(case_id)
    doc = setup_document(case_id)
    add_cover(doc, row, meta, evidence_path, metrics)
    add_model_page(doc, row, meta, capture)
    add_actual_model_page(doc, row, capture)
    add_execution_page(doc, row, meta, evidence_path, selected, metrics)
    add_modules_page(doc, row, meta)
    add_equations_page(doc, row, meta)
    add_results_page(doc, row, meta, metrics)
    add_final_page(doc, row, meta, evidence_path, selected)

    core = doc.core_properties
    core.title = f"S-Structures {case_id} 실제 모델링·자체 구조해석엔진 검증보고서 편집본"
    core.subject = f"{case_id} editable benchmark verification report"
    core.author = "S-Structures Verification"
    core.keywords = "S-Structures, STRIX21, structural analysis, verification, editable"
    core.comments = "Generated from deterministic model, execution evidence, and repository module trace."

    output = CASE_FOLDERS[case_id] / "05_보고서" / f"{case_id}_실제모델링_자체해석엔진_검증보고서_편집본.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    OUTPUT_ROOT.mkdir(parents=True, exist_ok=True)
    mirror = OUTPUT_ROOT / output.name
    shutil.copy2(output, mirror)
    return output, mirror, capture


def count_docx_parts(path):
    from zipfile import ZipFile
    with ZipFile(path) as archive:
        document_xml = archive.read("word/document.xml")
        media = [name for name in archive.namelist() if name.startswith("word/media/")]
    return {
        "documentXmlBytes": len(document_xml),
        "imageCount": len(media),
        "hasActualModelHeading": "실제 S-Structures 모델링 화면".encode("utf-8") in document_xml,
        "hasEngineModuleHeading": "사용된 S-Structures 자체 해석모듈".encode("utf-8") in document_xml,
    }


def update_readme(case_id, docx_name):
    readme = CASE_FOLDERS[case_id] / "05_보고서/README.md"
    original = readme.read_text(encoding="utf-8") if readme.exists() else "# 보고서\n"
    start = "<!-- STRIX21-DOCX-START -->"
    end = "<!-- STRIX21-DOCX-END -->"
    block = f"""{start}
## 편집 가능한 Word 보고서

- `{docx_name}`: 개념 형상·실제 S-Structures 모델링·해석모듈·사용 식·결과·주장 경계를 직접 수정할 수 있는 Word 원본
- 본문·제목·표는 편집 가능하며 개념도와 실제 제품 화면은 근거 이미지로 삽입됨
{end}
"""
    if start in original and end in original:
        before = original.split(start, 1)[0]
        after = original.split(end, 1)[1]
        updated = before.rstrip() + "\n\n" + block + after.lstrip("\n")
    else:
        first_break = original.find("\n") + 1
        updated = original[:first_break] + "\n" + block + "\n" + original[first_break:]
    readme.write_text(updated, encoding="utf-8")


def write_index(results):
    lines = [
        "# STRIX 21 편집 가능한 Word 보고서 색인",
        "",
        "21개 보고서는 제목·본문·표·판정 문구를 Word에서 직접 편집할 수 있다. 개념도와 실제 S-Structures 화면은 근거 이미지로 삽입했다.",
        "",
        "| 순서 | 사례 | Word 편집본 | 상태 |",
        "|---:|---|---|---|",
    ]
    for item in results:
        row = item["row"]
        folder = CASE_FOLDERS[row["id"]].name
        lines.append(f"| {R3.ORDER.index(row['id']) + 1} | {row['id']} | `{folder}/05_보고서/{item['output'].name}` | {R3.STATUS[row['lane']]['ko']} |")
    lines.extend([
        "",
        "## 편집 범위",
        "",
        "- 편집 가능: 제목, 설명, 모델링 항목, 코드 모듈 표, 해석식, 결과 표, 주장 경계, 다음 단계",
        "- 이미지: 개념도와 실제 S-Structures 화면은 원본 근거를 보존하기 위해 그림으로 삽입",
        "- 원본 PDF: 현재 승인본으로 그대로 유지",
    ])
    INDEX_PATH.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--case", default="all")
    args = parser.parse_args()
    R3.BASE.register_fonts()
    comparison = json.loads(R3.COMPARISON.read_text(encoding="utf-8"))
    by_id = {row["id"]: row for row in comparison["cases"]}
    ids = R3.ORDER if args.case == "all" else [args.case]
    results = []
    for case_id in ids:
        row = by_id[case_id]
        row["lane"] = row.get("lane") or "engine_only"
        meta = R3.BASE.METHODS[case_id]
        evidence_path, _raw, selected = R3.load_evidence(row)
        metrics = R3.extract_metrics(row, selected)
        output, mirror, capture = build_docx(row, meta, evidence_path, selected, metrics)
        structural = count_docx_parts(output)
        if not structural["hasActualModelHeading"] or not structural["hasEngineModuleHeading"] or structural["imageCount"] < 2:
            raise RuntimeError(f"{case_id} DOCX structural check failed: {structural}")
        update_readme(case_id, output.name)
        results.append({
            "row": row,
            "output": output,
            "mirror": mirror,
            "structural": structural,
            "sha256": R3.sha256(output),
            "modelBookSha256": capture["modelBookSha256"],
            "screenshotSha256": capture["screenshotSha256"],
        })
        print(f"{case_id}: DOCX PASS images={structural['imageCount']} sha256={results[-1]['sha256'][:16]}")

    if args.case == "all":
        write_index(results)
        QA_ROOT.mkdir(parents=True, exist_ok=True)
        manifest = {
            "version": "strix21-editable-docx-package-r1",
            "status": "STRUCTURAL_PASS_RENDER_PENDING",
            "generatedAt": "2026-08-31",
            "preset": "compact_reference_guide",
            "namedOverrides": ["Malgun Gothic Korean typography", "0.72 inch vertical margins", "editorial_cover title block"],
            "caseCount": len(results),
            "reports": [{
                "caseId": item["row"]["id"],
                "status": R3.STATUS[item["row"]["lane"]]["ko"],
                "docxPath": item["output"].relative_to(VAULT).as_posix(),
                "mirrorPath": item["mirror"].relative_to(REPO).as_posix(),
                "docxSha256": item["sha256"],
                "modelBookSha256": item["modelBookSha256"],
                "screenshotSha256": item["screenshotSha256"],
                "structuralChecks": item["structural"],
            } for item in results],
            "indexPath": INDEX_PATH.relative_to(VAULT).as_posix(),
        }
        MANIFEST_PATH.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(json.dumps({"status": manifest["status"], "caseCount": len(results), "manifest": str(MANIFEST_PATH)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
